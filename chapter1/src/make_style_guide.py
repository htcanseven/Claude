"""Build House_Style.docx from House_Style.md, in the house style it describes.

The document doubles as the Word template for contributors: it defines the
Caption paragraph style and the HS Book Table table style, its own tables are
set in that table style, and it ends with an example figure and caption taken
from Chapter 1.

usage:  python3 src/make_style_guide.py      (run from chapter1/)
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, Emu, RGBColor

sys.path.insert(0, 'src')
from apply_house_style import add_styles, house_table, child, ppr, rpr, q, TEXT_EMU  # noqa: E402

SRC, DST = 'House_Style.md', 'House_Style.docx'
TNR = 'Times New Roman'
INK = RGBColor(0x1a, 0x1a, 0x1a)
EXAMPLE_FIG = 'figures/ch1/fig_1_05.png'
EXAMPLE_CAP = ('Figure 1.5.', ' Loss components against rotational speed at fixed machine geometry: '
               'copper loss approximately constant, iron loss rising roughly with the square of '
               'speed and windage with its cube. The loss values are illustrative.')


def runs(par, text, size=None, font=None):
    """Inline markup: **bold**, *italic*, `code`, <sub>x</sub>."""
    for tok in re.split(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`|<sub>.+?</sub>)', text):
        if not tok:
            continue
        if tok.startswith('**'):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*'):
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith('`'):
            r = par.add_run(tok[1:-1]); r.font.name = 'Courier New'; r.font.size = Pt(10)
            continue
        elif tok.startswith('<sub>'):
            r = par.add_run(tok[5:-6]); r.font.subscript = True
        else:
            r = par.add_run(tok)
        if size:
            r.font.size = size
        if font:
            r.font.name = font


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Cm(2.54))
    st = doc.styles['Normal']
    st.font.name, st.font.size = TNR, Pt(11)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.08
    for name, size in (('Title', 18), ('Heading 1', 16), ('Heading 2', 13)):
        s = doc.styles[name]
        s.font.name, s.font.size, s.font.bold, s.font.color.rgb = TNR, Pt(size), True, INK
        rf = s.element.rPr.find(q('rFonts'))              # theme bindings would win over
        for k in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):   # the name
            rf.attrib.pop(q(k), None)
        s.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
        s.paragraph_format.space_after = Pt(6)
    for name in ('List Bullet', 'List Number'):
        doc.styles[name].font.name = TNR


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r':?-{3,}:?', c) for c in r)]
    ncol = len(cells[0])
    t = doc.add_table(rows=len(cells), cols=ncol)
    for i, row in enumerate(cells):
        for j, v in enumerate(row[:ncol]):
            p = t.rows[i].cells[j].paragraphs[0]
            runs(p, v)
    weight = [min(60, max(8, max(len(re.sub(r'[*`]|</?sub>', '', r[j])) for r in cells)))
              for j in range(ncol)]
    for g, w in zip(t._tbl.tblGrid.findall(q('gridCol')), weight):
        g.set(q('w'), str(w * 100))               # house_table rescales to the full width
    house_table(t._tbl, 'L' * ncol, head=bool(cells[0][0] or cells[0][1]), label='guide')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    doc = Document()
    setup(doc)
    lines = open(SRC, encoding='utf-8').read().split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith('```'):
            j = i + 1
            while not lines[j].strip().startswith('```'):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.6)
                r = p.add_run(lines[j]); r.font.name = 'Courier New'; r.font.size = Pt(9.5)
                j += 1
            doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
            i = j + 1
            continue
        if s.startswith('|'):
            j = i
            while j < len(lines) and lines[j].strip().startswith('|'):
                j += 1
            table(doc, lines[i:j])
            i = j
            continue
        if s.startswith('# '):
            doc.add_paragraph(s[2:], style='Title')
        elif s.startswith('## '):
            doc.add_paragraph(s[3:], style='Heading 1')
        elif s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); runs(p, s[2:])
        elif re.match(r'^\d+\. ', s):
            p = doc.add_paragraph(style='List Number'); runs(p, re.sub(r'^\d+\. ', '', s))
        elif s:
            p = doc.add_paragraph(); runs(p, s)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1

    # an example figure and caption, exactly as they stand in Chapter 1
    doc.add_paragraph('Example: a figure and its caption as they stand in Chapter 1',
                      style='Heading 1')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(EXAMPLE_FIG, width=Emu(TEXT_EMU))
    cap = doc.add_paragraph()
    child(ppr(cap._p), 'pStyle', {'val': 'Caption'})
    r = cap.add_run(EXAMPLE_CAP[0]); r.bold = True
    cap.add_run(EXAMPLE_CAP[1])

    add_styles(doc.styles.element)
    zoom = doc.settings.element.find(q('zoom'))   # python-docx's template omits the
    if zoom is not None:                           # percentage the schema requires
        zoom.set(q('percent'), '100')
    doc.save(DST)
    print(f'wrote {DST}')


if __name__ == '__main__':
    main()
