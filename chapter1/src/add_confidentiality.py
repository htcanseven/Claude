"""Stamp a confidentiality notice onto the two documents that go out to
invited contributors.

The proposal has not been submitted to Wiley, so the annotated table of
contents and the sample chapter are unpublished drafts.  A sentence in the
covering email protects nothing once the file is forwarded on its own, so the
notice is put into the documents themselves.

It deliberately permits one kind of onward sharing.  The book is organised so
that a chapter owner recruits their own co-authors, which they cannot do
without showing them the plan; a blanket prohibition would be broken by the
very people we are asking to recruit.

usage:  python3 src/add_confidentiality.py        (run from chapter1/)
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

YEAR = 2026
AUTHORS = 'H. T. Canseven, I. Petrov and J. Pyrhönen'
LEAD = 'Confidential — unpublished draft.  '

BODY = ('{what} is an unpublished working document prepared for a book proposal that has '
        'not yet been submitted to Wiley-IEEE Press. It is sent to you in confidence and '
        'solely so that you can judge whether to take part. Please do not forward, '
        'circulate, post or quote it, in whole or in part, without the authors’ '
        'written permission. If you invite a co-author, you may of course show it to them '
        'on the same terms. ')
RIGHTS = f'© {YEAR} {AUTHORS}. All rights reserved.'

JOBS = [
    ('incoming/Annotated_Table_of_Contents_v2.docx', 'to_send/Annotated_Table_of_Contents.docx',
     'This annotated table of contents', 4),   # after the title block
    ('incoming/Chapter_1_final.docx', 'to_send/Chapter_1.docx',
     'This draft chapter', 0),                 # the chapter opens straight into §1.1
]


def bottom_rule(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '6'), ('w:color', '999999')):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    # w:pBdr has a fixed position in CT_PPr: after w:pStyle, before w:spacing.
    style = pPr.find(qn('w:pStyle'))
    pPr.insert(list(pPr).index(style) + 1 if style is not None else 0, bdr)


def stamp(src, dst, what, after):
    doc = Document(src)
    note = doc.add_paragraph()                      # appended, then moved into place
    anchor = doc.paragraphs[after]
    anchor._p.addprevious(note._p)

    for text, bold in ((LEAD, True), (BODY.format(what=what), False), (RIGHTS, False)):
        run = note.add_run(text)
        run.bold = bold
        run.italic = not bold
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    note.paragraph_format.space_after = Pt(14)
    bottom_rule(note)

    os.makedirs(os.path.dirname(dst), exist_ok=True)
    doc.save(dst)
    return dst


if __name__ == '__main__':
    for src, dst, what, after in JOBS:
        out = stamp(src, dst, what, after)
        print(f'{out}  {os.path.getsize(out) / 1024:.0f} kB   ({what.lower()})')
