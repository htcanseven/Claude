"""Re-emit Figure 1.2 as native PowerPoint shapes.

The drawing is neither rasterised nor embedded as a picture: every rectangle,
polygon, centre line and label built by `src/make_schematics.py` is walked out
of the matplotlib figure and written again as an editable PowerPoint shape, so
the figure can be moved, recoloured and relabelled in PowerPoint without going
back to Python.  Section hatching becomes a PowerPoint pattern fill, dash-dot
centre lines become dash-dot connectors, and every text label stays text.

usage:  python3 src/mpl_to_pptx.py          (run from chapter1/)
"""
import runpy
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.colors as mcolors
from matplotlib.patches import Rectangle, Polygon, FancyBboxPatch, Circle
from matplotlib.text import Annotation, Text

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_PATTERN, MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

OUT = 'figures/editable/Figure_1_2.pptx'
PANELS = ('(a)', '(b)')
SLIDE_W_IN = 10.0                 # the 6.6 in figure is enlarged to a 10 in slide
FONT = 'Times New Roman'
HATCH = {'///': MSO_PATTERN.LIGHT_UPWARD_DIAGONAL,
         '\\\\\\': MSO_PATTERN.LIGHT_DOWNWARD_DIAGONAL,
         'xx': MSO_PATTERN.DIAGONAL_CROSS,
         'x': MSO_PATTERN.DIAGONAL_CROSS}


# ─────────────────────────────────────────────────────────── colour helpers
def rgb(c):
    return RGBColor.from_string(mcolors.to_hex(c)[1:].upper())


def opaque(rgba):
    """True if the colour is actually painted."""
    return rgba is not None and len(rgba) == 4 and rgba[3] > 0.001


def set_alpha(fill, alpha):
    """python-pptx has no transparency API; write the a:alpha child directly."""
    if alpha >= 0.999:
        return
    for tag in ('a:srgbClr', 'a:schemeClr'):
        for clr in fill._xPr.iter(qn(tag)):
            a = clr.makeelement(qn('a:alpha'), {'val': str(int(round(alpha * 100000)))})
            clr.append(a)


# ───────────────────────────────────────────────────── the figure to convert
def load_panels():
    """Run make_schematics.py and pick the two axes of Figure 1.2 out of it."""
    ns = runpy.run_path('src/make_schematics.py')
    ax1, ax2 = ns['ax1'], ns['ax2']
    fig = ax1.figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    FigureCanvasAgg(fig)                 # plt.close() took the old canvas away
    fig.canvas.draw()
    return fig, (ax1, ax2), fig.canvas.get_renderer()


class Sheet:
    """Maps matplotlib display pixels onto the slide."""

    def __init__(self, fig, shapes):
        self.dpi = fig.dpi
        self.w_in, self.h_in = fig.get_size_inches()
        self.sc = SLIDE_W_IN / self.w_in
        self.shapes = shapes

    def px(self, ax, x, y):
        return ax.transData.transform((x, y))

    def emu(self, p):
        """display pixels (origin bottom-left) → slide EMU (origin top-left)"""
        return (Inches(p[0] / self.dpi * self.sc),
                Inches((self.h_in - p[1] / self.dpi) * self.sc))

    def pt(self, points):
        return Pt(points * self.sc)

    # ── primitives ────────────────────────────────────────────────────────
    def rect(self, p0, p1, rounded=None):
        (x0, y0), (x1, y1) = self.emu(p0), self.emu(p1)
        left, top = min(x0, x1), min(y0, y1)
        w, h = abs(x1 - x0), abs(y1 - y0)
        shape = self.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Emu(int(left)), Emu(int(top)), Emu(max(int(w), 1)), Emu(max(int(h), 1)))
        shape.shadow.inherit = False
        return shape

    def poly(self, pts, close=True):
        pts = [self.emu(p) for p in pts]
        b = self.shapes.build_freeform(int(pts[0][0]), int(pts[0][1]), scale=1.0)
        b.add_line_segments([(int(x), int(y)) for x, y in pts[1:]], close=close)
        shape = b.convert_to_shape()
        shape.shadow.inherit = False
        return shape

    def line(self, p0, p1):
        (x0, y0), (x1, y1) = self.emu(p0), self.emu(p1)
        return self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                         Emu(int(x0)), Emu(int(y0)),
                                         Emu(int(x1)), Emu(int(y1)))


# ────────────────────────────────────────────────────────────── styling
def style_fill(shape, fc, hatch, ec):
    if not opaque(fc):
        shape.fill.background()
        return
    if hatch and hatch in HATCH:
        shape.fill.patterned()
        shape.fill.pattern = HATCH[hatch]
        shape.fill.fore_color.rgb = rgb(ec[:3]) if opaque(ec) else RGBColor(0x1A, 0x1A, 0x1A)
        shape.fill.back_color.rgb = rgb(fc[:3])
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fc[:3])
    set_alpha(shape.fill, fc[3])


def style_line(shape, ec, lw, sheet, dash=None):
    ln = shape.line
    if not opaque(ec) or lw <= 0:
        ln.fill.background()
        return
    ln.color.rgb = rgb(ec[:3])
    ln.width = sheet.pt(lw)
    if dash is not None:
        ln._get_or_add_ln().set('cap', 'flat')
        ln.dash_style = dash
    set_alpha(ln.fill, ec[3])


def dash_of(artist):
    """Translate a matplotlib line style into the nearest PowerPoint dash."""
    try:
        offset, onoff = artist._unscaled_dash_pattern
    except AttributeError:
        offset, onoff = artist.get_linestyle(), None
    if not onoff:
        return None
    if len(onoff) >= 4:
        return MSO_LINE_DASH_STYLE.DASH_DOT
    if len(onoff) == 2:
        return MSO_LINE_DASH_STYLE.DASH if onoff[0] > 2.5 else MSO_LINE_DASH_STYLE.SQUARE_DOT
    return None


# ────────────────────────────────────────────────────────────── emitters
def name_of(p):
    """A name for the PowerPoint selection pane, from what the patch is."""
    kind = ('rounded box' if isinstance(p, FancyBboxPatch) else
            'polygon' if isinstance(p, Polygon) else
            'circle' if isinstance(p, Circle) else 'box')
    h = p.get_hatch()
    return f'sectioned {kind}' if h else kind


def emit_patch(sheet, ax, p):
    fc, ec = p.get_facecolor(), p.get_edgecolor()
    lw, hatch = p.get_linewidth(), p.get_hatch()

    if isinstance(p, FancyBboxPatch):
        (x0, y0), (x1, y1) = p.get_path().get_extents().get_points()
        shape = sheet.rect(sheet.px(ax, x0, y0), sheet.px(ax, x1, y1), rounded=True)
        # rounding_size is in data units; take it through the axes transform
        r = getattr(p.get_boxstyle(), 'rounding_size', 0.0) or 0.0
        rr = abs(sheet.px(ax, x0 + r, y0)[0] - sheet.px(ax, x0, y0)[0]) / sheet.dpi * sheet.sc
        shape.adjustments[0] = min(0.5, rr / max(min(shape.width, shape.height) / 914400.0, 1e-6))
    elif isinstance(p, Rectangle):
        x0, y0 = p.get_x(), p.get_y()
        shape = sheet.rect(sheet.px(ax, x0, y0),
                           sheet.px(ax, x0 + p.get_width(), y0 + p.get_height()))
    elif isinstance(p, Polygon):
        xy = p.get_xy()
        if len(xy) > 1 and np.allclose(xy[0], xy[-1]):
            xy = xy[:-1]
        shape = sheet.poly([sheet.px(ax, x, y) for x, y in xy], close=True)
    elif isinstance(p, Circle):
        cx, cy = p.center
        r = p.radius
        shape = sheet.rect(sheet.px(ax, cx - r, cy - r), sheet.px(ax, cx + r, cy + r))
        shape._element.find(qn('p:spPr')).find(qn('a:prstGeom')).set('prst', 'ellipse')
    else:
        return None
    style_fill(shape, fc, hatch, ec)
    style_line(shape, ec, lw, sheet)
    return shape


def emit_line(sheet, ax, ln):
    xs, ys = ln.get_xdata(), ln.get_ydata()
    if len(xs) < 2:
        return
    col = mcolors.to_rgba(ln.get_color(), ln.get_alpha())
    pts = [sheet.px(ax, x, y) for x, y in zip(xs, ys)]
    if len(pts) == 2:
        shape = sheet.line(pts[0], pts[1])
    else:
        shape = sheet.poly(pts, close=False)
        shape.fill.background()
    style_line(shape, col, ln.get_linewidth(), sheet, dash_of(ln))
    return shape


def emit_text(sheet, ax, t):
    s = t.get_text()
    if not s.strip():
        return
    lines = s.split('\n')
    fs = t.get_fontsize()
    x, y = sheet.px(ax, *t.get_position()) if t.get_transform() is ax.transData \
        else t.get_transform().transform(t.get_position())
    lead = getattr(t, '_linespacing', 1.2)
    lead = float(lead) if isinstance(lead, (int, float)) else 1.2
    h_in = len(lines) * fs * lead / 72.0 * sheet.sc
    w_in = max(len(l) for l in lines) * fs * 0.64 / 72.0 * sheet.sc + 0.05

    ha, va = t.get_horizontalalignment(), t.get_verticalalignment()
    desc = 0.21 * fs / 72.0 * sheet.sc          # PowerPoint anchors on the descender,
    left = {'center': -w_in / 2, 'left': 0.0, 'right': -w_in}[ha]   # matplotlib on the baseline
    top = {'center': -h_in / 2, 'top': 0.0, 'bottom': -h_in,
           'baseline': -h_in + desc, 'center_baseline': -h_in / 2}[va]

    ox, oy = sheet.emu((x, y))
    box = sheet.shapes.add_textbox(Emu(int(ox + Inches(left))), Emu(int(oy + Inches(top))),
                                   Inches(w_in), Inches(h_in))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'center': MSO_ANCHOR.MIDDLE, 'top': MSO_ANCHOR.TOP}.get(
        va, MSO_ANCHOR.BOTTOM)
    align = {'center': PP_ALIGN.CENTER, 'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT}[ha]
    col = mcolors.to_rgba(t.get_color(), t.get_alpha())

    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        # An exact spacing, not a percentage: PowerPoint reckons a percentage
        # against the font's own line height, matplotlib against the point size.
        para.line_spacing = sheet.pt(fs * lead)
        run = para.add_run()
        run.text = line
        f = run.font
        f.name, f.size = FONT, sheet.pt(fs)
        f.bold = t.get_fontweight() in ('bold', 'heavy', 'semibold', 600, 700, 800, 900)
        f.italic = t.get_style() == 'italic'
        f.color.rgb = rgb(col[:3])
    return box


def emit_leader(sheet, ax, ann, renderer):
    """The line an annotate() draws from its label to the point it names."""
    props = getattr(ann, 'arrowprops', None)
    if not props:
        return
    p1 = np.array(sheet.px(ax, *ann.xy), dtype=float)          # the point named
    p0 = np.array(sheet.px(ax, *ann.get_position()), dtype=float)   # the label
    scale = sheet.dpi / 72.0
    if ann.get_text().strip():                       # start at the edge of the label
        # Text.get_window_extent, not Annotation's: the latter unions in the arrow.
        bb = Text.get_window_extent(ann, renderer).expanded(1.06, 1.20)
        pad = props.get('shrinkA', 2) * scale
        ts = np.linspace(0, 1, 400)
        inside = [t for t in ts if bb.contains(*(p0 + t * (p1 - p0)))]
        if inside:
            p0 = p0 + min(1.0, max(inside) + pad / max(np.linalg.norm(p1 - p0), 1e-6)) * (p1 - p0)
    d = p1 - p0
    n = np.linalg.norm(d)
    if n < 1e-6:
        return
    p1 = p1 - d / n * props.get('shrinkB', 2) * scale

    shape = sheet.line(p0, p1)
    col = mcolors.to_rgba(props.get('color', '0.4'))
    style_line(shape, col, props.get('lw', 1.0), sheet)
    style = props.get('arrowstyle', '-')
    ln = shape.line._get_or_add_ln()
    if '<' in style:
        ln.append(ln.makeelement(qn('a:headEnd'), {'type': 'triangle', 'w': 'sm', 'len': 'sm'}))
    if '>' in style:
        ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'sm', 'len': 'sm'}))
    return shape


# ──────────────────────────────────────────────────────────────── the walk
def convert():
    fig, panels, renderer = load_panels()

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_W_IN * fig.get_size_inches()[1] / fig.get_size_inches()[0])
    slide = prs.slides.add_slide(prs.slide_layouts[6])          # blank
    sheet = Sheet(fig, slide.shapes)

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    bg.line.fill.background(); bg.shadow.inherit = False
    bg.name = 'background'

    counts = {'patch': 0, 'line': 0, 'text': 0, 'leader': 0}
    skipped = []
    for panel, ax in zip(PANELS, panels):
        # Each panel is its own group, so it can be moved or scaled as one part
        # and its pieces are still reachable by double-clicking into the group.
        group = slide.shapes.add_group_shape()
        group.name = f'Figure 1.2 {panel}'
        sheet.shapes = group.shapes
        n = 0

        # PowerPoint stacks in insertion order, matplotlib in zorder: sort first.
        items = ([(a.get_zorder(), i, 'patch', a) for i, a in enumerate(ax.patches)] +
                 [(a.get_zorder(), i, 'line', a) for i, a in enumerate(ax.lines)] +
                 [(a.get_zorder(), i, 'text', a) for i, a in enumerate(ax.texts)])
        items.sort(key=lambda it: (it[0], it[2] == 'text', it[1]))
        for _, _, kind, a in items:
            if kind == 'patch':
                shape = emit_patch(sheet, ax, a)
                if shape is None:
                    skipped.append(type(a).__name__)
                    continue
                n += 1
                shape.name = f'{panel} {name_of(a)} {n:02d}'
                counts['patch'] += 1
            elif kind == 'line':
                shape = emit_line(sheet, ax, a)
                n += 1
                if shape is not None:
                    shape.name = f'{panel} centre line {n:02d}'
                counts['line'] += 1
            else:
                if isinstance(a, Annotation):
                    shape = emit_leader(sheet, ax, a, renderer)
                    n += 1
                    if shape is not None:
                        shape.name = f'{panel} leader {n:02d}'
                    counts['leader'] += 1
                box = emit_text(sheet, ax, a)
                if box is not None:
                    n += 1
                    box.name = f'{panel} label “{a.get_text().splitlines()[0][:28]}”'
                    counts['text'] += 1

    sheet.shapes = slide.shapes
    slide.notes_slide.notes_text_frame.text = (
        'Figure 1.2 of Chapter 1, drawn as native PowerPoint shapes — no picture is '
        'embedded, so every part, hatch and label can be edited here.\n\n'
        'Source of truth: chapter1/src/make_schematics.py. Regenerate this file with '
        '`python3 src/mpl_to_pptx.py`; the PNG that goes into the manuscript comes from '
        'make_schematics.py itself, so edits made here do not travel back.\n\n'
        'Section hatching is a PowerPoint pattern fill: `///` and `\\\\\\` on adjacent '
        'parts, cross-hatch on the housing castings.')
    prs.save(OUT)
    print(f'{OUT}: {sum(counts.values())} shapes on one '
          f'{prs.slide_width.inches:.1f} × {prs.slide_height.inches:.1f} in slide')
    print('   ' + '  '.join(f'{k} {v}' for k, v in counts.items()))
    if skipped:
        print('   NOT converted:', ', '.join(sorted(set(skipped))))


if __name__ == '__main__':
    convert()
